"""网关测试夹具：假框架目录（真 pytest 跑假 test_xlsx）、假结果库、假串口控制台、假服务端客户端。"""

from __future__ import annotations

import json
import os
import shutil
import sys
import textwrap
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

# gateway/vendor/cex_core 不入库（带内部模板身份，本仓守零内部资产），由人手跑
# tools/sync_gateway_vendor.py --only cex_core 从 skills 仓生成。测试对仓库只读：
# 从不改写 gateway/vendor（找到哪棵 skills 树就拿它覆盖一遍，会把半截改动或别的分支
# 悄悄带进打包与别的工具）；
# 与 skills 仓是否一致由 test_vendor_readonly.py 用 --check 报。没生成过就不收集网关测试。
SKILLS_ROOT = Path(os.environ.get("CEX_SKILLS_ROOT") or REPO_ROOT.parent / "compile-excel-skills")
SYNC_COMMAND = (f"python3 tools/sync_gateway_vendor.py --only cex_core "
                f"--skills-root {SKILLS_ROOT}")
VENDOR_NOTE = ""
if not (REPO_ROOT / "gateway" / "vendor" / "cex_core" / "__init__.py").is_file():
    VENDOR_NOTE = ("gateway tests not collected: gateway/vendor/cex_core was never generated; "
                   f"run {SYNC_COMMAND}")
    collect_ignore_glob = ["test_*.py"]


def pytest_report_header(config):
    return VENDOR_NOTE or None

TEMPLATE = REPO_ROOT / "gateway" / "vendor" / "cex_core" / "templates" / "case_template.xlsx"
BUILD = "SAMPLE_BUILD_LOCAL"
GRAMMAR = {
    "destructive_commands": {"patterns": [r"^clear\s+config\s+all\b", r"\breboot\b"]},
    "bed_probes": {"build": {"cmd": "show version", "extract": r"Software Version\s*:\s*(.+)"}},
}
CRED_LITERAL = "Zq9SecretLiteral"

FAKE_TEST_XLSX = '''
import json, os, time
from pathlib import Path

from openpyxl import load_workbook


def test_case(request):
    here = Path(__file__).parent
    root = Path(os.getcwd())
    build = request.config.getoption("--build")
    slow = root / "slow_run"
    if slow.exists():
        time.sleep(float(slow.read_text() or "3"))
    verdicts = json.loads((root / "fake_verdicts.json").read_text()) \\
        if (root / "fake_verdicts.json").exists() else {}
    wb = load_workbook(here / "case.xlsx", read_only=True)
    ids = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            a = str(row[0]).strip() if row and row[0] is not None else ""
            if a.isdigit() and len(a) >= 12 and a != "999999999999999" and a not in ids:
                ids.append(a)
    module = here.parent.name[len("ist_staging_"):]
    # 与真框架一致：每次运行一个报告目录，结果库一行一案，url 指向本次运行的报告目录
    run = "run-%d-%s" % (time.time_ns(), build)
    rows = json.loads((root / "fake_results.json").read_text()) \\
        if (root / "fake_results.json").exists() else []
    for cid in ids:
        if cid in verdicts.get("skip", []):
            continue
        rel = "report/%s/%s/ist_staging_%s/%s/test_xlsx/case.xlsx/%s" % (run, module, module, here.name, cid)
        (root / rel).mkdir(parents=True, exist_ok=True)
        result = str(verdicts.get(cid, "PASS")).upper()
        # 与真框架一样收尾：计数、PASS/FAIL 横幅，紧跟 end case（unfinished 里的案模拟停在案里）
        close = "" if cid in verdicts.get("unfinished", []) else (
            "#\\n################# The failed check point num:   %d   ####\\n#\\n"
            "################# The passed check point num:   %d   ####\\n#\\n"
            "######################      %s      ####################\\n#######   end case: %s\\n"
            % (0 if result == "PASS" else 1, 1 if result == "PASS" else 0, result, cid))
        (root / rel / (cid + ".txt")).write_text("ran %s on %s in %s\\n%s" % (cid, build, run, close))
        (root / rel / "apv_192.0.2.10.txt").write_text("APV(config)#show sdns node\\nnode1 %s\\n" % cid)
        sub = "ist_staging_" + module
        rows = [r for r in rows if not (r["table"] == build and r["case_id"] == cid and r["sub_module"] == sub)]
        rows.append({"table": build, "case_id": cid, "sub_module": sub,
                     "result": verdicts.get(cid, "PASS"),
                     "url": "http://jumphost/test/fw/" + rel + "/" + cid + "/"})
    (root / "fake_results.json").write_text(json.dumps(rows))
'''

FAKE_CONFTEST = '''
def pytest_addoption(parser):
    parser.addoption("--build", action="store", default="")
'''

FAKE_MYSQLDB = '''
import json, os, re


class Result_DB(object):
    def db_exec(self, queries):
        sql, params = queries[0]
        table = re.search(r"FROM `([^`]+)`", sql).group(1)
        rows = json.load(open(os.path.join(os.getcwd(), "fake_results.json")))
        wanted = [p for p in params if not p.endswith("%")]
        return [(r["case_id"], r["result"], r["url"]) for r in rows
                if r["table"] == table and r["case_id"] in wanted]
'''

FAKE_CU = r'''
import os, sys, time
hang_on = os.environ.get("FAKE_CU_HANG_ON", "")
out = sys.stdout
out.write("Connected.\r\n"); out.flush()
state = "start"
for raw in sys.stdin:
    line = raw.strip()
    if state == "start":
        out.write("login: "); state = "user"
    elif state == "user":
        out.write("Password: "); state = "pass"
    elif state == "pass":
        out.write("APV#"); state = "enable"
    elif line in ("conf ter", "config ter"):
        out.write("APV(config)#"); state = "config"
    elif state == "config":
        if hang_on and line == hang_on:
            out.flush(); time.sleep(30)
        out.write("APV(config)#")
    else:
        out.write("APV#")
    out.flush()
'''


class FakeServer:
    """鸭子类型的 ServerClient：令牌表 + 规则文件。"""

    def __init__(self):
        self.tokens = {
            "run-token": {"active": True, "username": "alice", "scope": "jumphost:run"},
            "admin-token": {"active": True, "username": "root",
                            "scope": "jumphost:run jumphost:admin"},
            "bob-token": {"active": True, "username": "bob", "scope": "jumphost:run"},
        }
        self.grammar = GRAMMAR

    def introspect(self, token):
        return self.tokens.get(token, {"active": False})

    def fetch_bundle_file(self, build, path):
        from gateway.introspect import IntrospectError

        if self.grammar is None:
            raise IntrospectError("server unreachable (test)")
        return json.dumps(self.grammar).encode()


@pytest.fixture()
def fake_env(tmp_path, monkeypatch):
    apv = tmp_path / "apv_src"
    (apv / "lib").mkdir(parents=True)
    (apv / "conf").mkdir()
    (apv / "smoke_test" / "sdns").mkdir(parents=True)
    (apv / "conftest.py").write_text(FAKE_CONFTEST, encoding="utf-8")
    (apv / "lib" / "test_xlsx.py").write_text(FAKE_TEST_XLSX, encoding="utf-8")
    (apv / "lib" / "mysqldb.py").write_text(FAKE_MYSQLDB, encoding="utf-8")
    (apv / "lib" / "__init__.py").write_text("", encoding="utf-8")
    (apv / "lib" / "excel_contract.json").write_text("{}", encoding="utf-8")
    (apv / "lib" / "creds.py").write_text(f'password = "{CRED_LITERAL}"\n', encoding="utf-8")
    (apv / "conf" / "bed.conf").write_text(textwrap.dedent("""\
        [comm]
        ssh_ips = 127.0.0.1, 127.0.0.2
        ports = eth1, eth2, eth3
        [other]
        mysql_ip = 127.0.0.1
        [array_ustack]
        user = admin
        passwd = devpass
        hostname = APV
        """), encoding="utf-8")
    fake_cu = tmp_path / "fake_cu.py"
    fake_cu.write_text(FAKE_CU, encoding="utf-8")
    secret = tmp_path / "client.secret"
    secret.write_text("s", encoding="utf-8")
    secret.chmod(0o600)
    config = tmp_path / "gateway.toml"
    config.write_text(textwrap.dedent(f"""\
        [server]
        url = "http://127.0.0.1:1"
        client_id = "gateway"
        client_secret_file = "{secret}"
        build = "{BUILD}"
        [framework]
        apv_src = "{apv}"
        py38 = "{sys.executable}"
        conf_name = "bed"
        staging_parent = "{apv / 'smoke_test' / 'sdns'}"
        default_module = "sdns"
        run_max_s = 60
        [state]
        dir = "{tmp_path / 'state'}"
        lease_ttl_s = 600
        [device]
        console_command = ["{sys.executable}", "{fake_cu}", "{{tty}}"]
        [init_device]
        commands = ["no page", "clear config all", "ip add {{port1}} 192.0.2.7{{idx}} 24"]
        long_commands = {{ "clear config all" = 5 }}
        step_timeout_s = 3
        login_timeout_s = 5
        """), encoding="utf-8")
    from gateway.config import load
    from gateway.tools import Gateway

    server = FakeServer()
    gateway = Gateway(load(config), server=server)
    return {"gateway": gateway, "server": server, "apv": apv, "config": config,
            "tmp": tmp_path}


def make_workbook(path: Path, rows: list[tuple[str, str, str, str]]) -> bytes:
    """rows: (autoid 或空, E, F, G)。以冻结模板为底写执行页。"""
    from openpyxl import load_workbook

    from gateway.vendor.cex_core.ist_emit.excel_contract import resolve_execution_sheet

    shutil.copy2(TEMPLATE, path)
    wb = load_workbook(path)
    ws, _ = resolve_execution_sheet(wb, allow_legacy=False)
    start = ws.max_row + 1
    for offset, (autoid, device, method, command) in enumerate(rows):
        row = start + offset
        if autoid:
            ws.cell(row=row, column=1, value=autoid)
        ws.cell(row=row, column=5, value=device)
        ws.cell(row=row, column=6, value=method)
        ws.cell(row=row, column=7, value=command)
    wb.save(path)
    return path.read_bytes()
