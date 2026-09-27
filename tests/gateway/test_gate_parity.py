"""上机前闸与框架看到的是同一份东西：公式缓存值、999999999999999 之后的行、G 拆参去引号后的实参。

框架（lib/test_xlsx.py）用 data_only=True 读缓存值、跑到文件末尾、把 G 拆开去引号再交给设备；
闸若只看公式原文、在哨兵处停、只查 G 原文，`clear config all` 就能换个写法上机。
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pytest

from conftest import CRED_LITERAL, GRAMMAR, TEMPLATE, make_workbook

from gateway import gates

A1, A2 = "202609260000000001", "202609260000000002"
SENTINEL = gates.SENTINEL_AUTOID
LITERALS = frozenset({CRED_LITERAL})


def _check(data: bytes) -> gates.FrozenCase:
    frozen = gates.freeze(data)
    gates.check(frozen, grammar=GRAMMAR, credential_literals=LITERALS)
    return frozen


def _problems(data: bytes) -> str:
    with pytest.raises(gates.GateError) as exc:
        _check(data)
    return "\n".join(exc.value.problems)


def _with_cached_formula(path: Path, sheet_xml: str, ref: str, formula: str, cached: str) -> bytes:
    """把 ref 格改成带缓存值的公式（Excel 存盘后的样子：<f> 公式 + <v> 缓存值）。"""
    out = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == sheet_xml:
                text = data.decode("utf-8")
                cell = re.compile(r'<c r="%s"[^>]*?(?:/>|>.*?</c>)' % ref)
                assert cell.search(text), f"{ref} not in {sheet_xml}"
                text = cell.sub(f'<c r="{ref}" t="str"><f>{formula}</f><v>{cached}</v></c>',
                                text, count=1)
                data = text.encode("utf-8")
            dst.writestr(info, data)
    return out.getvalue()


def _execution_row() -> int:
    """make_workbook 追加的第一行的行号（模板执行页之后）。"""
    from openpyxl import load_workbook

    from gateway.vendor.cex_core.ist_emit.excel_contract import resolve_execution_sheet

    ws, _ = resolve_execution_sheet(load_workbook(TEMPLATE), allow_legacy=False)
    return ws.max_row + 1


def test_formula_whose_cached_value_is_destructive_is_refused(tmp_path):
    path = tmp_path / "f.xlsx"
    make_workbook(path, [(A1, "APV_0", "cmd_config", "show version")])
    row = _execution_row()
    data = _with_cached_formula(path, "xl/worksheets/sheet1.xml", f"G{row}",
                                "&quot;show &quot;&amp;&quot;version&quot;", "clear config all")
    text = _problems(data)
    assert f"执行!G{row}: holds a formula" in text


def test_value_starting_with_equals_is_refused(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "eq.xlsx"
    make_workbook(path, [(A1, "APV_0", "cmd_config", "show version")])
    row = _execution_row()
    wb = load_workbook(path)
    cell = wb["执行"].cell(row=row, column=4)
    cell.value = "=== teardown ==="
    cell.data_type = "s"           # 存成字符串而不是公式
    wb.save(path)
    assert f"执行!D{row}: holds a formula" in _problems(path.read_bytes())


def test_credential_literal_hidden_in_a_formula_cache_on_another_sheet(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "notes.xlsx"
    make_workbook(path, [(A1, "APV_0", "cmd", "show version")])
    wb = load_workbook(path)
    notes = wb.create_sheet("notes")
    notes["B2"] = '="x"'
    wb.save(path)
    clean = _with_cached_formula(path, "xl/worksheets/sheet2.xml", "B2", "&quot;x&quot;", "x")
    assert _check(clean).autoids == (A1,), "别的页的公式不参与执行，缓存值干净就放行"
    leaked = _with_cached_formula(path, "xl/worksheets/sheet2.xml", "B2",
                                  "&quot;pw &quot;&amp;&quot;Zq9&quot;", f"pw {CRED_LITERAL}")
    text = _problems(leaked)
    assert "notes!B2: contains a framework credential literal" in text
    assert CRED_LITERAL not in text


def test_rows_after_the_sentinel_are_gated_and_their_cases_counted(tmp_path):
    bad = make_workbook(tmp_path / "s1.xlsx", [
        (A1, "APV_0", "cmd", "show version"), (SENTINEL, "", "", ""),
        (A2, "APV_0", "cmd_config", "clear config all"), (SENTINEL, "", "", "")])
    assert "'clear config all' matches destructive rule" in _problems(bad)
    good = make_workbook(tmp_path / "s2.xlsx", [
        (A1, "APV_0", "cmd", "show version"), (SENTINEL, "", "", ""),
        (A2, "APV_0", "cmd", "show slb real"), (SENTINEL, "", "", "")])
    assert _check(good).autoids == (A1, A2), "哨兵不算用例，它后面的用例照算"


@pytest.mark.parametrize("command", [
    '"clear config all"', "' clear config all '", "cmd=clear config all",
    "cmd='clear config all', timeout=60", 'cmd = "clear config all"'])
def test_g_is_checked_as_the_framework_passes_it_to_the_device(tmp_path, command):
    data = make_workbook(tmp_path / "q.xlsx", [(A1, "APV_0", "cmd_config", command)])
    text = _problems(data)
    # 报的是这一格：直接给出框架交给设备的实参，或给原文并注明实际执行的串（新版扫描器带 executed）
    assert re.search(r"执行!G\d+: (?:'clear config all'|.+\(runs as 'clear config all'\)) "
                     r"matches destructive rule", text), text
    assert text.count("matches destructive rule") == 1, "同一格同一条规则只报一次"


def test_ordinary_quoted_and_keyword_arguments_still_pass(tmp_path):
    data = make_workbook(tmp_path / "ok.xlsx", [
        (A1, "APV_0", "cmd_config", 'slb real http "r1" 10.0.0.1 80'),
        ("", "APV_0", "cmd", "cmd=show version, timeout=30"),
        ("", "APV_0", "cmds_config", "slb group method g1 rr\nslb group member g1 r1"),
        ("", "check_point", "found", "reboot")])   # 断言期望文本不交给设备
    assert _check(data).autoids == (A1,)


def test_framework_arguments_mirror_the_framework_split():
    fa = gates.framework_arguments
    assert fa('"clear config all"', "cmd_config") == ["clear config all"]
    assert fa("cmd=clear config all", "cmd_config") == ["clear config all"]
    assert fa('a, b="c, d", \'e\'', "cmd") == ["a", "c, d", "e"]
    assert fa("x\\,y", "cmd") == ["x\\,y"], "转义的逗号不拆"
    assert fa("a,b", "cmds_config") == ["a,b"], "cmds_config/execute 整格一个实参"
    assert fa("l1\nl2", "cmd") == ["l1\nl2"], "多行格整格一个实参"
    assert fa('"unclosed', "cmd") == ['"unclosed'], "拆不开框架整卷拒跑，这里只回原文"
