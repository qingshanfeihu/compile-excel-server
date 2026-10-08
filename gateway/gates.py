"""上机前闸：上传的工作簿先在内存里冻结，再逐项检查；任何一项不过就拒绝，不改写文件。

对应 InfoTest dev_run_batch 身份链的“冻结 + 派生”几步（批次工具第 2–4 步）：
- zip 结构与体积（成员数、单成员/总大小、压缩比、路径穿越、重复成员）；
- openpyxl 能在行列上限内打开，执行页唯一（契约 marker 与表头行）；
- 执行页 autoid 从冻结字节里读，不信调用方给的列表；
另加两道 InfoTest 上机路径上没有的闸：
- 自毁命令（规则来自服务端数据包的 domain_grammar.json，读不到就拒绝）；
- 凭据字面量（从跳板机上真实框架源码 AST 提取，任何单元格命中就拒绝；报告只给单元格位置）。

闸查的必须就是框架执行的（框架 lib/test_xlsx.py）：
- 框架用 data_only=True 读公式的缓存值，闸读不到缓存值 → 执行页有公式（或以 = 开头的值）就拒收；
- 框架一直跑到文件末尾，999999999999999 之后的行照样执行 → 每一行都查（它本身不算用例）；
- 框架把 G 按逗号拆开、去引号、取关键字参数的值再交给设备 → 这些实参与原文一起查。
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .vendor.cex_core.ist_emit._sealed_io import validate_xlsx_zip_budget
from .vendor.cex_core.ist_emit.excel_contract import resolve_execution_sheet
from .vendor.cex_core.scan_destructive import (
    DestructiveRulesUnavailable,
    load_patterns,
    scan_lines,
)

MAX_XLSX_BYTES = 32 * 1024 * 1024
MAX_ROWS = 20000
MAX_COLUMNS = 64
SENTINEL_AUTOID = "999999999999999"
_AUTOID_RE = re.compile(r"^\d{12,24}$")
_MIN_LITERAL_LEN = 4
_MAX_LISTED = 20


class GateError(ValueError):
    """工作簿被拒；problems 是逐项原因（英文，给模型看）。"""

    def __init__(self, problems: list[str]):
        super().__init__("; ".join(problems))
        self.problems = problems


@dataclass(frozen=True)
class FrozenCase:
    data: bytes
    sha256: str
    size: int
    autoids: tuple[str, ...]
    command_lines: tuple[tuple[str, str], ...]
    cells: tuple[tuple[str, str], ...]
    # 用到的被测设备对象（E 列 APV_k / Segk_tmp）→ 第一次出现的单元格；框架按 conf
    # [comm] ssh_ips 的第 k 个地址连它，conf 里没有第 k 台就整卷跑不起来
    devices: tuple[tuple[str, str], ...] = ()


# 框架 lib/test_xlsx.py 的设备表：APV_0/1/2 走 apv_xlsx(…, k)，Seg0/1/2_tmp 走 conftest 里
# ssh_ips[k] 的 segment 夹具
_DEVICE_OBJECT_RE = re.compile(r"^(?:APV_(\d)|Seg(\d)_tmp)$")


def device_index(obj: str) -> int | None:
    """E 列对象要用 conf 里第几台设备；不是被测设备对象就是 None。"""
    match = _DEVICE_OBJECT_RE.match(obj)
    if match is None:
        return None
    return int(match.group(1) or match.group(2))


# ── 框架怎样把 G 交给设备（lib/test_xlsx.py 的 _split_parameter_parts / _unquote_parameter /
#    _keyword_split / _raw_call_arguments，逐行照抄语义，不 import 框架）────────────────
def _split_parameter_parts(text: str) -> list[str]:
    parts: list[str] = []
    current: list[str] = []
    quote: str | None = None
    escaped = False
    for char in text:
        if escaped:
            current.append(char)
            escaped = False
            continue
        if char == "\\":
            current.append(char)
            escaped = True
            continue
        if quote is not None:
            current.append(char)
            if char == quote:
                quote = None
            continue
        if char in ('"', "'"):
            current.append(char)
            quote = char
            continue
        if char == ",":
            part = "".join(current).strip()
            if part:
                parts.append(part)
            current = []
            continue
        current.append(char)
    if quote is not None:
        raise ValueError("unclosed quote in G")
    tail = "".join(current).strip()
    if tail:
        parts.append(tail)
    return parts


def _unquote_parameter(value: str) -> str:
    text = value.strip()
    if len(text) >= 2 and text[0] in ('"', "'") and text[-1] == text[0]:
        return text[1:-1].strip()
    return text


def _keyword_split(part: str) -> tuple[str, str] | None:
    quote: str | None = None
    escaped = False
    for index, char in enumerate(part):
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if quote is not None:
            if char == quote:
                quote = None
            continue
        if char in ('"', "'"):
            quote = char
            continue
        if char == "=":
            key = part[:index].strip()
            if re.match(r"^[A-Za-z_]\w*$", key):
                return key, part[index + 1:].strip()
            return None
    return None


def framework_arguments(text: str, method: str) -> list[str]:
    """框架把这一格 G 交给设备方法的每个字符串实参：execute/cmds_config 与多行格整格一个；
    其余按不在引号里的逗号拆开，位置参数去引号，关键字参数（cmd=...）取值去引号。
    拆不开（引号未闭合）框架整卷拒跑，这里只回原文。"""
    if method in ("execute", "cmds_config") or "\n" in text or "\r" in text:
        return [text]
    try:
        parts = _split_parameter_parts(text)
    except ValueError:
        return [text]
    values = []
    for part in parts:
        keyword = _keyword_split(part)
        values.append(_unquote_parameter(keyword[1] if keyword else part))
    return values


def _cached_values(data: bytes, where: list[tuple[str, str]]) -> list[tuple[str, str]]:
    """执行页以外的公式格：取框架与读表人看到的缓存值，只交给凭据闸。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    try:
        out = []
        for title, coordinate in where:
            value = wb[title][coordinate].value
            if isinstance(value, str) and value.strip():
                out.append((f"{title}!{coordinate}", value))
        return out
    finally:
        wb.close()


def freeze(data: bytes) -> FrozenCase:
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise GateError(["workbook is empty"])
    data = bytes(data)
    if len(data) > MAX_XLSX_BYTES:
        raise GateError([f"workbook exceeds {MAX_XLSX_BYTES} bytes"])
    try:
        validate_xlsx_zip_budget(data, error_type=GateError, message="workbook zip is malformed")
    except GateError:
        raise GateError(["workbook zip is malformed or exceeds size/ratio limits"]) from None
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = [info.filename for info in archive.infolist()]
        if len(names) != len(set(names)):
            raise GateError(["workbook zip has duplicate members"])
        for info in archive.infolist():
            name = info.filename
            if name.startswith("/") or ".." in name.split("/") or "\\" in name:
                raise GateError(["workbook zip has a member path outside the archive"])
            if info.flag_bits & 0x1:
                raise GateError(["workbook zip has encrypted members"])
            if (info.external_attr >> 16) & 0o170000 == 0o120000:
                raise GateError(["workbook zip has symlink members"])
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=False, data_only=False)
    except Exception as exc:  # noqa: BLE001 — 任何解析失败都是拒收理由
        raise GateError([f"workbook does not open: {type(exc).__name__}"]) from None
    try:
        for ws in wb.worksheets:
            if ws.max_row > MAX_ROWS or ws.max_column > MAX_COLUMNS:
                raise GateError([f"sheet {ws.title!r} exceeds {MAX_ROWS} rows or "
                                 f"{MAX_COLUMNS} columns"])
        try:
            ws, _ = resolve_execution_sheet(wb, allow_legacy=False)
        except Exception as exc:  # noqa: BLE001
            raise GateError([f"execution sheet does not satisfy the Excel contract: {exc}"]) \
                from None
        formulas: list[str] = []
        elsewhere: list[tuple[str, str]] = []
        cells: list[tuple[str, str]] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    where = f"{sheet.title}!{cell.coordinate}"
                    if sheet is ws and (cell.data_type == "f" or (
                            isinstance(value, str) and value.startswith("="))):
                        formulas.append(where)
                    elif cell.data_type == "f":
                        elsewhere.append((sheet.title, cell.coordinate))
                    if isinstance(value, str) and value.strip():
                        cells.append((where, value))
        if formulas:
            raise GateError([f"{where}: holds a formula; the framework runs its cached value, "
                             "which this gate cannot check - write the literal value"
                             for where in formulas[:_MAX_LISTED]]
                            + ([f"... and {len(formulas) - _MAX_LISTED} more formula cells"]
                               if len(formulas) > _MAX_LISTED else []))
        autoids: list[str] = []
        commands: list[tuple[str, str]] = []
        devices: dict[str, str] = {}
        # 不在 999999999999999 处停：框架一直跑到文件末尾，它后面的行照样执行
        for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
            first = str(row[0]).strip() if row and row[0] is not None else ""
            if _AUTOID_RE.match(first) and first != SENTINEL_AUTOID and first not in autoids:
                autoids.append(first)
            device = str(row[4] or "").strip() if len(row) > 4 else ""
            if device_index(device) is not None:
                devices.setdefault(device, f"{ws.title}!E{row_no}")
            method = str(row[5] or "").strip() if len(row) > 5 else ""
            raw = row[6] if len(row) > 6 else None
            if not device.startswith("APV") or raw is None or not str(raw).strip():
                continue
            where, text = f"{ws.title}!G{row_no}", str(raw)
            arguments = framework_arguments(text, method)
            cells.extend((where, value) for value in arguments if value != text)
            seen: set[str] = set()
            for value in (text, *arguments):
                for line in value.splitlines():
                    if line.strip() and line not in seen:
                        seen.add(line)
                        commands.append((where, line))
    finally:
        wb.close()
    if elsewhere:
        cells.extend(_cached_values(data, elsewhere))
    if not autoids:
        raise GateError(["execution sheet has no case autoids"])
    return FrozenCase(data=data, sha256=hashlib.sha256(data).hexdigest(), size=len(data),
                      autoids=tuple(autoids), command_lines=tuple(commands), cells=tuple(cells),
                      devices=tuple(devices.items()))


def check(frozen: FrozenCase, *, grammar: dict[str, Any] | Path | None,
          credential_literals: frozenset[str], device_count: int | None = None) -> None:
    """device_count：本床框架 conf 里的设备台数（[comm] ssh_ips）；给了就拒收用到第 k 台
    （k ≥ 台数）设备对象的工作簿——框架连那台时取不到地址，整卷一个案都不跑。"""
    problems: list[str] = []
    if device_count is not None:
        for obj, where in frozen.devices:
            index = device_index(obj)
            if index is not None and index >= device_count:
                problems.append(
                    f"{where}: {obj} needs device {index}, but this bed's framework conf lists "
                    f"{device_count} device(s) (APV_0..APV_{device_count - 1}); the framework "
                    "could not reach it and no case in the workbook would run")
    try:
        patterns = load_patterns(grammar if grammar is not None else {})
    except DestructiveRulesUnavailable as exc:
        raise GateError([f"destructive-command rules unavailable, refusing to run: {exc}"]) \
            from None
    reported: set[tuple[str, str]] = set()
    for finding in scan_lines(frozen.command_lines, patterns):
        # 同一格的原文与拆出的实参命中同一条规则只报一次
        if (finding["where"], finding["rule"]) in reported:
            continue
        reported.add((finding["where"], finding["rule"]))
        runs_as = f" (runs as {finding['executed']!r})" if finding.get("executed") else ""
        problems.append(f"{finding['where']}: {finding['command']!r}{runs_as} matches destructive "
                        f"rule {finding['rule']!r}; clean up only what the case created")
    literals = [value for value in credential_literals if len(value) >= _MIN_LITERAL_LEN]
    flagged: set[str] = set()
    for where, text in frozen.cells:
        folded = text.casefold()
        if where not in flagged and any(value.casefold() in folded for value in literals):
            flagged.add(where)
            problems.append(f"{where}: contains a framework credential literal; "
                            "remove it (credentials come from the framework conf at run time)")
    if problems:
        raise GateError(problems)
